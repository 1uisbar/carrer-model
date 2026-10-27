#!/usr/bin/env python3
"""Genera modelo_carrera.xlsx con fórmulas vivas a partir de supuestos.json.

Hojas:
  Supuestos      selectores (ubicación, vivienda, escenario), parámetros efectivos y TODAS las variables
  Ubic_<id>      una por ubicación: parámetros + cotización de vivienda por configuración (fórmulas)
  <camino>       una por camino: tabla de etapas (editable) + simulación anual 2027-2046 con fórmulas
  Comparativa    patrimonio por año de todos los caminos, compra/pago de vivienda, cruces

Las fechas de las etapas vienen del motor (scripts/export_paths.mjs) con los supuestos por defecto;
son celdas azules editables. Si cambias fechas en supuestos.json, regenera con `npm run excel`.
"""
import json
import os
import re
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = json.load(open(os.path.join(ROOT, 'supuestos.json')))
subprocess.run(['node', os.path.join(ROOT, 'scripts', 'export_paths.mjs')], check=True, cwd=ROOT, stdout=subprocess.DEVNULL)
D = json.load(open(os.path.join(ROOT, 'excel', 'paths_default.json')))

START, H = A['general']['anioInicio']['v'], A['general']['horizonte']['v']
YEARS = list(range(START, START + H))
LOCS = list(A['ubicaciones'].keys())
CFGS = list(A['vivienda']['configuracion']['opciones'].keys())
ESCS = ['pesimista', 'base', 'optimista']
MAXPH, MAXEV = 25, 20

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
INPUT = Font(color='1F4FD8')
HDR_FILL = PatternFill('solid', fgColor='E8ECF5')
NOTE = Font(italic=True, color='5B6477')

wb = Workbook()
names = {}


def add_name(name, ref):
    names[name] = ref
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def absref(ws, cell):
    col, row = re.match(r'([A-Z]+)(\d+)', cell).groups()
    return f"'{ws.title}'!${col}${row}"


def nm(path):
    return path.replace('.', '_')


def hdr(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = BOLD
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical='top')


# ----------------------------------------------------------------------------- Supuestos
S = wb.active
S.title = 'Supuestos'
S['A1'] = 'Supuestos del modelo (editar celdas azules)'
S['A1'].font = TITLE
S['A2'] = 'Fuente única: supuestos.json. Cambiar aquí recalcula todo el libro. Porcentajes como fracción (0.07 = 7%). Booleanos 1/0.'
S['A2'].font = NOTE

S['A3'], S['B3'] = 'Ubicación (id)', 'dfw'
S['A4'], S['B4'] = 'Configuración de vivienda', A['vivienda']['configuracion']['v']
S['A5'], S['B5'] = 'Escenario', 'base'
for c in ('B3', 'B4', 'B5'):
    S[c].font = INPUT
for i, l in enumerate(LOCS):
    S.cell(row=3 + i, column=8, value=l)
    S.cell(row=3 + i, column=9, value=A['ubicaciones'][l]['nombre'])
for i, c in enumerate(CFGS):
    S.cell(row=3 + i, column=10, value=c)
    S.cell(row=3 + i, column=11, value=A['configVivienda'][c]['unidadesRentadas'])
for i, e in enumerate(ESCS):
    S.cell(row=3 + i, column=12, value=e)
S['H2'], S['J2'], S['K2'], S['L2'] = 'Ubicaciones', 'Configuraciones', 'Unid. rentadas', 'Escenarios'
dv = DataValidation(type='list', formula1=f'$H$3:$H${2 + len(LOCS)}'); S.add_data_validation(dv); dv.add('B3')
dv = DataValidation(type='list', formula1=f'$J$3:$J${2 + len(CFGS)}'); S.add_data_validation(dv); dv.add('B4')
dv = DataValidation(type='list', formula1=f'$L$3:$L$5'); S.add_data_validation(dv); dv.add('B5')
S['C3'] = f'=MATCH(B3,$H$3:$H${2 + len(LOCS)},0)'
S['C4'] = f'=MATCH(B4,$J$3:$J${2 + len(CFGS)},0)'
S['C5'] = '=MATCH(B5,$L$3:$L$5,0)'
add_name('sel_loc', absref(S, 'B3')); add_name('sel_locIdx', absref(S, 'C3'))
add_name('sel_cfg', absref(S, 'B4')); add_name('sel_cfgIdx', absref(S, 'C4'))
add_name('sel_escIdx', absref(S, 'C5'))
add_name('cfg_units', absref(S, 'K3').replace('$K$3', f'$K$3:$K${2 + len(CFGS)}'))

# escenarios (tabla editable)
row = 10
S.cell(row=row - 1, column=1, value='Ajustes por escenario').font = BOLD
hdr(S, row, ['Ajuste'] + ESCS)
for i, k in enumerate(A['escenarios']['base'].keys()):
    r = row + 1 + i
    S.cell(row=r, column=1, value=k)
    for j, e in enumerate(ESCS):
        c = S.cell(row=r, column=2 + j, value=A['escenarios'][e][k]); c.font = INPUT
    S.cell(row=r, column=5, value=f'=CHOOSE(sel_escIdx,B{r},C{r},D{r})')
    add_name('esc_' + k, absref(S, f'E{r}'))
S.cell(row=row, column=5, value='Activo').font = BOLD
row = row + 2 + len(A['escenarios']['base'])

# lista de todas las hojas-valor (escalares)
leaf_rows_start = row + 32
S.cell(row=leaf_rows_start - 1, column=1, value='Todas las variables').font = TITLE
hdr(S, leaf_rows_start, ['Ruta', 'Descripción', 'Valor', 'Unidad', 'Fuente', 'URL', 'Confianza', 'Estimado', 'Nota'])
r = leaf_rows_start + 1
TABLES = {}


def walk(node, path):
    global r
    if isinstance(node, dict) and 'v' in node:
        v = node['v']
        if isinstance(v, list):
            TABLES[path] = (node, v)
            return
        if path.startswith('ubicaciones.'):
            return
        if isinstance(v, bool):
            v = 1 if v else 0
        f = A['fuentes'].get(node.get('f', ''), {})
        S.cell(row=r, column=1, value=path)
        S.cell(row=r, column=2, value=node.get('l'))
        c = S.cell(row=r, column=3, value=v); c.font = INPUT
        if node.get('u') == '%':
            c.number_format = '0.00%'
        S.cell(row=r, column=4, value=node.get('u'))
        S.cell(row=r, column=5, value=f.get('t'))
        S.cell(row=r, column=6, value=f.get('url'))
        S.cell(row=r, column=7, value=node.get('c'))
        S.cell(row=r, column=8, value='sí' if node.get('e') else '')
        S.cell(row=r, column=9, value=node.get('n'))
        if path != 'vivienda.configuracion':
            add_name(nm(path), absref(S, f'C{r}'))
        r += 1
        return
    if isinstance(node, dict):
        for k, ch in node.items():
            if k in ('fuentes', '_meta', 'escenarios', 'configVivienda'):
                continue
            walk(ch, f'{path}.{k}' if path else k)


walk(A, '')
for col, w in zip('ABCDEFGHI', (34, 52, 14, 8, 50, 40, 10, 9, 60)):
    S.column_dimensions[col].width = w

# parámetros efectivos (con escenario) y ubicación activa
S.cell(row=row, column=1, value='Parámetros efectivos (escenario aplicado)').font = BOLD
eff = [
    ('g_inf', 'Inflación efectiva', '=general_inflacion+esc_inflacion'),
    ('g_ret', 'Retorno efectivo', '=general_retornoInversion+esc_retorno'),
    ('g_cash', 'Retorno efectivo HYSA', '=general_retornoEfectivo'),
    ('g_appr', 'Apreciación efectiva', '=vivienda_apreciacion+esc_apreciacion'),
    ('g_mortRate', 'Tasa hipotecaria efectiva', '=MAX(0.01,vivienda_tasaHipoteca+esc_tasaHipoteca)'),
    ('g_salMult', 'Multiplicador salarios', '=esc_salarios'),
    ('g_gastos', 'Multiplicador gastos', '=esc_gastos'),
]
for i, (n, l, f) in enumerate(eff):
    rr = row + 1 + i
    S.cell(row=rr, column=1, value=l); S.cell(row=rr, column=2, value=n); S.cell(row=rr, column=3, value=f)
    add_name(n, absref(S, f'C{rr}'))

LOC_PARAMS = ['factorRN', 'factorMD', 'residentePGY1', 'gastosVida', 'rentaMensual', 'seguroPct', 'impuestoPropPct', 'factorMatch', 'rentaUnidad']
PRICE_KEYS = ['duplex', 'inlaw', 'dosCasas', 'casaAdu', 'aduCosto', 'multi']
TAXCODE = {'none': 1, 'oh': 2, 'ny': 3}

# tablas de brackets (forma diferencial: umbral inferior y cambio de tasa)
def bracket_table(key, col0):
    node, br = TABLES[key]
    c0 = col0
    S.cell(row=row, column=c0, value=node['l']).font = BOLD
    hdr(S, row + 1, ['Desde', 'Tasa', 'Δ tasa'], col=c0)
    lo = 0; prev = 0
    for i, (upto, rate) in enumerate(br):
        rr = row + 2 + i
        S.cell(row=rr, column=c0, value=lo).font = INPUT
        S.cell(row=rr, column=c0 + 1, value=rate).font = INPUT
        S.cell(row=rr, column=c0 + 2, value=f'={L(c0 + 1)}{rr}-{L(c0 + 1)}{rr - 1}' if i else f'={L(c0 + 1)}{rr}')
        lo = upto
    first, last = row + 2, row + 1 + len(br)
    add_name('tb_' + key.split('.')[-1] + '_lo', f"'Supuestos'!${L(c0)}${first}:${L(c0)}${last}")
    add_name('tb_' + key.split('.')[-1] + '_dr', f"'Supuestos'!${L(c0 + 2)}${first}:${L(c0 + 2)}${last}")


bracket_table('impuestos.federalBrackets', 5)
bracket_table('impuestos.nyBrackets', 9)
bracket_table('impuestos.nycBrackets', 13)

# ubicación activa (CHOOSE sobre las hojas de ubicación)
loc_row = row + 10
S.cell(row=loc_row - 1, column=1, value='Ubicación activa').font = BOLD
for i, p in enumerate(LOC_PARAMS + ['precio_' + k for k in PRICE_KEYS] + ['taxcode', 'nyc']):
    rr = loc_row + i
    S.cell(row=rr, column=1, value=p)
    S.cell(row=rr, column=3, value='=CHOOSE(sel_locIdx,' + ','.join(f'ubic_{l}_{p}' for l in LOCS) + ')')
    add_name('loc_' + p, absref(S, f'C{rr}'))
assert loc_row + len(LOC_PARAMS) + len(PRICE_KEYS) + 2 < leaf_rows_start - 1, 'reservar más filas'

# ----------------------------------------------------------------------------- hojas por ubicación
for l in LOCS:
    U = A['ubicaciones'][l]
    ws = wb.create_sheet('Ubic_' + l)
    ws['A1'] = U['nombre']; ws['A1'].font = TITLE
    ws['A2'] = U['notas']; ws['A2'].font = NOTE
    ws['A3'] = 'Licencia compacta: ' + U['compacta']; ws['A3'].font = NOTE
    hdr(ws, 5, ['Parámetro', 'Valor', 'Unidad', 'Fuente', 'Confianza', 'Estimado'])
    rr = 6
    items = [(p, U[p]) for p in LOC_PARAMS] + [('precio_' + k, U['precios'][k]) for k in PRICE_KEYS]
    for p, node in items:
        f = A['fuentes'].get(node.get('f', ''), {})
        ws.cell(row=rr, column=1, value=node['l'])
        c = ws.cell(row=rr, column=2, value=node['v']); c.font = INPUT
        if node.get('u') == '%':
            c.number_format = '0.00%'
        ws.cell(row=rr, column=3, value=node.get('u'))
        ws.cell(row=rr, column=4, value=f.get('t'))
        ws.cell(row=rr, column=5, value=node.get('c'))
        ws.cell(row=rr, column=6, value='sí' if node.get('e') else '')
        add_name(f'ubic_{l}_{p}', absref(ws, f'B{rr}'))
        rr += 1
    for p, v, lbl in (('taxcode', TAXCODE[U['impuesto']], 'Impuesto estatal (1=ninguno, 2=Ohio, 3=Nueva York)'), ('nyc', 1 if U['nyc'] else 0, 'Impuesto municipal NYC (1/0)')):
        ws.cell(row=rr, column=1, value=lbl); c = ws.cell(row=rr, column=2, value=v); c.font = INPUT
        add_name(f'ubic_{l}_{p}', absref(ws, f'B{rr}')); rr += 1
    # cotización de vivienda 2026
    rr += 1
    ws.cell(row=rr, column=1, value='Cotización de vivienda (precios 2026, escenario activo)').font = BOLD
    rr += 1
    hdr(ws, rr, ['Configuración', 'Precio', 'Enganche', 'Cierre', 'Préstamo', 'P&I mensual', 'Impuesto/mes', 'Seguro/mes', 'PMI/mes', 'Mantenimiento/mes', 'Renta recibida/mes', 'PITI', 'Costo neto/mes', 'Ingreso bruto anual necesario'])
    for j, cfg in enumerate(CFGS):
        q = rr + 1 + j
        units = A['configVivienda'][cfg]['unidadesRentadas']
        ws.cell(row=q, column=1, value=A['vivienda']['configuracion']['opciones'][cfg])
        ws.cell(row=q, column=2, value=f'=ubic_{l}_precio_{cfg}')
        ws.cell(row=q, column=3, value=f'=B{q}*vivienda_enganchePct')
        ws.cell(row=q, column=4, value=f'=B{q}*vivienda_cierrePct')
        ws.cell(row=q, column=5, value=f'=B{q}-C{q}')
        ws.cell(row=q, column=6, value=f'=PMT(g_mortRate/12,vivienda_plazoHipoteca*12,-E{q})')
        ws.cell(row=q, column=7, value=f'=B{q}*ubic_{l}_impuestoPropPct/12')
        ws.cell(row=q, column=8, value=f'=B{q}*ubic_{l}_seguroPct/12')
        ws.cell(row=q, column=9, value=f'=IF(vivienda_enganchePct<0.2,E{q}*vivienda_pmiTasa/12,0)')
        ws.cell(row=q, column=10, value=f'=B{q}*vivienda_mantenimientoPct/12')
        ws.cell(row=q, column=11, value=f'={units}*ubic_{l}_rentaUnidad*(1-vivienda_vacancia)')
        ws.cell(row=q, column=12, value=f'=F{q}+G{q}+H{q}+I{q}')
        ws.cell(row=q, column=13, value=f'=L{q}+J{q}-K{q}*(1-vivienda_gestionPct)')
        ws.cell(row=q, column=14, value=f'=12*MAX(0,L{q}-0.75*K{q})/vivienda_dtiMax')
        for c in range(2, 15):
            ws.cell(row=q, column=c).number_format = '#,##0'
    ws.column_dimensions['A'].width = 46
    for c in 'BCDEFGHIJKLMN':
        ws.column_dimensions[c].width = 14

# ----------------------------------------------------------------------------- hojas por camino
ROLE_GROUP = {}
for rid, R in A['roles'].items():
    ROLE_GROUP[rid] = R['grupo']


def role_salary_expr(role):
    if role == 'cm':
        return 'personal_tarifaHora*2080*loc_factorRN/ubic_miami_factorRN'
    if role == 'residente':
        return 'loc_residentePGY1'
    if role == 'travel':
        return 'roles_travel_base*tiempos_travelSemanas'
    g = ROLE_GROUP[role]
    f = {'rn': 'loc_factorRN', 'remoto': '(1+(loc_factorRN-1)*0.5)', 'crna': '(1+(loc_factorRN-1)*0.5)', 'md': 'loc_factorMD'}.get(g, '1')
    return f'roles_{role}_base*{f}'


def raise_expr(role):
    if role == 'residente':
        return 'g_inf+0.005'
    return f'roles_{role}_aumento'


def mult_expr(ph):
    m = ph['mult']
    role = ph.get('role')
    c = A['tiempos']['cdiCertAumento']['v']
    if role == 'cm' and abs(m - (1 + A['personal']['bsnAumento']['v'])) < 1e-9:
        return '=1+personal_bsnAumento'
    if role == 'cdi':
        for n in (1, 2):
            if abs(m - (1 + c) ** n) < 1e-9:
                return f'=(1+tiempos_cdiCertAumento)^{n}'
    if role == 'residente':
        for n in (1, 2, 3):
            if abs(m - (1 + A['roles']['residente']['aumento']['v']) ** n) < 1e-9:
                return f'=(1+roles_residente_aumento)^{n}'
    return m


def hours_expr(ph):
    h = ph['hours']
    if ph.get('role') != 'cm':
        return h
    if ph.get('tuition'):
        return '=tiempos_npHorasTrabajo'
    if 'estudio' in ph['lbl'].lower():
        return '=md_horasTrabajoEstudiando/40'
    return '=personal_horasSemana/40'


def tuition_expr(ph):
    t = ph.get('tuition') or 0
    if not t:
        return 0
    cands = {
        'tiempos_crnaMatricula': A['tiempos']['crnaMatricula']['v'], 'tiempos_npMatricula': A['tiempos']['npMatricula']['v'],
        'tiempos_msnLiderazgoCosto': A['tiempos']['msnLiderazgoCosto']['v'],
        'md_usceCostoMes*12*md_usceDuracion': A['md']['usceCostoMes']['v'] * 12 * A['md']['usceDuracion']['v'],
    }
    for e, v in cands.items():
        if abs(v - t) < 1e-6:
            return '=' + e
    return t


def event_cost_expr(e):
    lbl, cost = e['lbl'], e.get('cost') or 0
    table = [
        ('Matrícula restante BSN', 'personal_bsnCostoRestante'), ('Step 1 reprobado', 'md_costoStep1'), ('Step 1', 'md_costoStep1+md_costoPrepUSMLE'),
        ('Step 2 CK reprobado', 'md_costoStep2'), ('Step 2 CK + ECFMG', 'md_costoStep2+md_costoECFMG'), ('ERAS', 'md_costoERAS+md_costoEntrevistas'),
        ('Mudanza', 'md_costoMudanza'), ('Step 3', 'md_costoLicencia'), ('Curso de inyecciones', 'tiempos_injectorCurso'),
    ]
    for prefix, expr in table:
        if lbl.startswith(prefix):
            return '=' + expr
    return cost


PH_COLS = ['Etapa', 'Inicio', 'Fin', 'Rol', 'Salario base FT 2026 ($)', 'Aumento anual', 'Horas (FTE)', 'Multiplicador', 'Bono al inicio',
           'Matrícula total (2026$)', 'Estudio (1/0)', 'PSLF (1/0)', 'Match 403b (1/0)', 'Travel (1/0)', '% stipend libre', 'S-corp anual base', 'Autoempleo anual base', 'Attending (1/0)']
OVC = 20  # columna T: matriz de traslape
EVC = OVC + H - 1
Y0 = 33   # primera fila de la tabla anual

# columnas de la simulación anual: (clave, encabezado, fórmula). Tokens: {X} = misma fila col X; {X-1} = fila anterior; {k}=índice año; {r}=fila; {OV}=columna de traslape
SIM = [
    ('year', 'Año', None), ('k', 'k', None), ('fI', 'Índice inflación', '=(1+g_inf)^({year}-2026)'),
    ('w2', 'Salarios W-2', '=SUMPRODUCT({OV},$E$5:$E$29*(1+$F$5:$F$29)^({year}-2026)*$G$5:$G$29*$H$5:$H$29*(1-$N$5:$N$29*$O$5:$O$29))*g_salMult+SUMPRODUCT(($B$5:$B$29>={year})*($B$5:$B$29<{year}+1)*$I$5:$I$29)*g_salMult'),
    ('w2m', 'W-2 con match', '=SUMPRODUCT({OV},$E$5:$E$29*(1+$F$5:$F$29)^({year}-2026)*$G$5:$G$29*$H$5:$H$29*$M$5:$M$29)*g_salMult'),
    ('stip', 'Stipends libres', '=SUMPRODUCT({OV},$E$5:$E$29*(1+$F$5:$F$29)^({year}-2026)*$G$5:$G$29*$H$5:$H$29*$N$5:$N$29*$O$5:$O$29)*g_salMult'),
    ('scorp', 'Ingreso S-corp', '=SUMPRODUCT({OV},$P$5:$P$29*(1+$F$5:$F$29)^({year}-2026))'),
    ('se', 'Autoempleo (LLC/tutoría)', '=SUMPRODUCT({OV},$Q$5:$Q$29*(1+$F$5:$F$29)^({year}-2026))+IF(personal_llcActivo=1,personal_llcNeto*(1+personal_llcCrecimiento)^{k},0)'),
    ('tuit', 'Matrícula', '=SUMPRODUCT({OV},$J$5:$J$29/($C$5:$C$29-$B$5:$B$29+($C$5:$C$29=$B$5:$B$29)))*{fI}'),
    ('study', 'Fracción estudio', '=SUMPRODUCT({OV},$K$5:$K$29)'),
    ('pslf', 'Fracción PSLF', '=SUMPRODUCT({OV},$L$5:$L$29)'),
    ('trav', 'Fracción travel', '=SUMPRODUCT({OV},$N$5:$N$29)'),
    ('att', 'Attending', '=IF(SUMPRODUCT({OV},$R$5:$R$29)>0,1,0)'),
    ('ev', 'Hitos (costos)', '=SUMPRODUCT(($AO$5:$AO$24>={year})*($AO$5:$AO$24<{year}+1)*$AP$5:$AP$24)*{fI}+IF({k}=0,IF(sel_loc<>"miami",personal_costoMudanzaInicial,0),0)'),
    ('k401', 'Aporte 403b', '=MIN(personal_aporte403bPct*{w2},retiro_limite403b*{fI})'),
    ('emp', 'Match empleador', '=MIN(personal_aporte403bPct,personal_match403bPct)*{w2m}'),
    ('sNet', 'S-corp neto', '=IF({scorp}>0,MAX(0,{scorp}-impuestos_scorpCostoAnual*{fI}),0)'),
    ('wS', 'Salario S-corp', '={sNet}*impuestos_scorpSalarioPct/(1+impuestos_ssTasa+impuestos_medicareTasa)'),
    ('eFS', 'FICA patronal S-corp', '=impuestos_ssTasa*MIN({wS},MAX(0,impuestos_ssTope*{fI}-{w2}))+impuestos_medicareTasa*{wS}'),
    ('dist', 'Distribución S-corp', '={sNet}-{wS}-{eFS}'),
    ('allW', 'Salarios totales', '={w2}+{wS}'),
    ('seB', 'Base autoempleo', '={se}*0.9235'),
    ('seT', 'Impuesto autoempleo', '=IF({se}>0,2*impuestos_ssTasa*MIN({seB},MAX(0,impuestos_ssTope*{fI}-{allW}))+2*impuestos_medicareTasa*{seB},0)'),
    ('fica', 'FICA empleado', '=impuestos_ssTasa*MIN({allW},impuestos_ssTope*{fI})+impuestos_medicareTasa*{allW}+impuestos_medicareAdicional*MAX(0,{allW}+{se}*0.9235-impuestos_medicareUmbral)'),
    ('agi', 'AGI', '=MAX(0,{w2}-{k401}+{wS}+{dist}+{se}-{seT}/2)'),
    ('ti0', 'Ingreso gravable antes QBI', '=MAX(0,{agi}-impuestos_deduccionEstandar*{fI})'),
    ('qbi', 'Deducción QBI', '=MIN(0.2*MAX(0,{dist}+{se}-{seT}/2)*MAX(0,MIN(1,1-({ti0}-impuestos_qbiUmbral*{fI})/(impuestos_qbiRango*{fI}))),0.2*{ti0})'),
    ('ti', 'Ingreso gravable', '=MAX(0,{ti0}-{qbi})'),
    ('fed', 'Impuesto federal', '=SUMPRODUCT(({ti}>tb_federalBrackets_lo*{fI})*({ti}-tb_federalBrackets_lo*{fI})*tb_federalBrackets_dr)'),
    ('nyTi', 'Gravable NY', '=MAX(0,{agi}-impuestos_nyDeduccion*{fI})'),
    ('st', 'Impuesto estatal/local', '=CHOOSE(loc_taxcode,0,impuestos_ohTasa*MAX(0,{agi}-impuestos_ohExento*{fI})+impuestos_ohMunicipal*{allW},SUMPRODUCT(({nyTi}>tb_nyBrackets_lo*{fI})*({nyTi}-tb_nyBrackets_lo*{fI})*tb_nyBrackets_dr)+loc_nyc*SUMPRODUCT(({nyTi}>tb_nycBrackets_lo*{fI})*({nyTi}-tb_nycBrackets_lo*{fI})*tb_nycBrackets_dr))'),
    ('tax', 'Impuestos totales', '={fed}+{st}+{fica}+{seT}'),
    ('cash', 'Ingreso personal bruto en efectivo', '={w2}+{stip}+{wS}+{dist}+{se}'),
    ('take', 'Efectivo neto (tras impuestos y 403b)', '={cash}-{tax}-{k401}'),
    ('gross', 'Ingreso bruto', '={w2}+{stip}+{scorp}+{se}'),
    ('liv', 'Gastos de vida', '=loc_gastosVida*{fI}*g_gastos'),
    # vivienda: cotización del año
    ('L0p', 'Ahorro líquido inicio', '=IF({k}=0,personal_liquidoInicial,{liq-1})'),
    ('own0', 'Ya era dueño', '=IF({k}=0,0,{own-1})'),
    ('basePr', 'Precio base 2026', '=CHOOSE(sel_cfgIdx,loc_precio_duplex,loc_precio_inlaw,loc_precio_dosCasas,loc_precio_casaAdu,loc_precio_multi)'),
    ('price', 'Precio en el año', '={basePr}*(1+g_appr)^({year}-2026)'),
    ('phys', 'Physician loan', '=IF(AND({att}=1,vivienda_prestamoMedico=1),1,0)'),
    ('dPct', '% enganche', '=IF({phys}=1,0.05,vivienda_enganchePct)'),
    ('loanQ', 'Préstamo', '={price}*(1-{dPct})'),
    ('pmtQ', 'Cuota P&I mensual', '=PMT(g_mortRate/12,vivienda_plazoHipoteca*12,-{loanQ})'),
    ('units', 'Unidades rentadas', '=INDEX(cfg_units,sel_cfgIdx)'),
    ('rentQ', 'Renta mensual (cotización)', '={units}*loc_rentaUnidad*(1+g_inf)^({year}-2026)*(1-vivienda_vacancia)'),
    ('pitiQ', 'PITI mensual', '={pmtQ}+{price}*loc_impuestoPropPct/12+{basePr}*loc_seguroPct*(1+vivienda_inflacionSeguro)^({year}-2026)/12+IF(AND({phys}=0,{dPct}<0.2),{loanQ}*vivienda_pmiTasa/12,0)'),
    ('need', 'Efectivo requerido', '={price}*{dPct}+{price}*vivienda_cierrePct+general_reservaMeses/12*{liv}'),
    ('othD', 'Otras deudas mensuales', '=(IF({k}=0,MIN(personal_irsSaldo,personal_irsPagoMensual*12),MIN({irsE-1},personal_irsPagoMensual*12))+IF({fedP}>0,PMT({fRP},MAX(1,prestamos_plazoAnios-{fYP}),-{fedP}),0)+IF({prvP}>0,PMT({pRP},MAX(1,prestamos_plazoAnios-{pYP}),-{prvP}),0))/12'),
    ('famC', 'Crédito aporte familiar', '=IF(OR(sel_cfg="duplex",sel_cfg="multi",sel_cfg="dosCasas"),0.75*personal_aporteFamiliaMensual,0)'),
    ('hM', 'Carga mensual para DTI', '={pitiQ}-0.75*{rentQ}-{famC}+{othD}'),
    ('incM', 'Ingreso mensual para DTI', '=({w2}+{scorp}+{se}+0.5*{stip})/12'),
    ('buy', 'Compra este año', '=IF(AND({own0}=0,{year}>=vivienda_anioMinCompra,{study}<0.5,{gross}>0,{L0p}>={need},{hM}<=vivienda_dtiMax*{incM}),1,0)'),
    ('own', 'Dueño', '=MAX({own0},{buy})'),
    ('buyY', 'Año de compra', '=IF({buy}=1,{year},IF({k}=0,"",{buyY-1}))'),
    ('L0', 'Ahorro tras enganche', '={L0p}-IF({buy}=1,{price}*{dPct}+{price}*vivienda_cierrePct,0)'),
    ('orig', 'Precio de compra', '=IF({buy}=1,{price},IF({k}=0,0,{orig-1}))'),
    ('mPmt', 'Cuota hipoteca', '=IF({buy}=1,{pmtQ},IF({k}=0,0,{mPmt-1}))'),
    ('pmiOn0', 'PMI activo (antes)', '=IF({buy}=1,IF(AND({phys}=0,{dPct}<0.2),1,0),IF({k}=0,0,{pmiOn-1}))'),
    ('isAdu', 'Config. con ADU', '=IF(sel_cfg="casaAdu",1,0)'),
    ('aduNow', 'Construye ADU', '=IF(AND({own}=1,{isAdu}=1,IF({k}=0,0,{aduB-1})=0,{year}>=IF({buyY}="",9999,{buyY})+vivienda_aduRetrasoAnios),1,0)'),
    ('aduB', 'ADU construido', '=MAX(IF({k}=0,0,{aduB-1}),{aduNow})'),
    ('aduPmt', 'Cuota ADU', '=IF({aduNow}=1,PMT((g_mortRate+0.005)/12,180,-loc_precio_aduCosto*{fI}),IF({k}=0,0,{aduPmt-1}))'),
    ('val0', 'Valor vivienda inicio', '=IF({buy}=1,{price},IF({k}=0,0,{valE-1}))+IF({aduNow}=1,loc_precio_aduCosto*{fI},0)'),
    ('mB', 'Hipoteca inicio', '=IF({buy}=1,{loanQ},IF({k}=0,0,{mE-1}))'),
    ('mRem', 'Hipoteca tras 12 pagos', '=IF({mB}<=0,0,FV(g_mortRate/12,12,{mPmt},-{mB}))'),
    ('mPaid', 'Pagado hipoteca', '=IF({mB}<=0,0,12*{mPmt}+MIN(0,{mRem}))'),
    ('mA', 'Hipoteca tras pagos', '=MAX(0,{mRem})'),
    ('aB', 'ADU inicio', '=IF({aduNow}=1,loc_precio_aduCosto*{fI},IF({k}=0,0,{aE-1}))'),
    ('aRem', 'ADU tras 12 pagos', '=IF({aB}<=0,0,FV((g_mortRate+0.005)/12,12,{aduPmt},-{aB}))'),
    ('aPaid', 'Pagado ADU', '=IF({aB}<=0,0,12*{aduPmt}+MIN(0,{aRem}))'),
    ('aA', 'ADU tras pagos', '=MAX(0,{aRem})'),
    ('ins', 'Seguro', '=IF({own}=1,{basePr}*loc_seguroPct*(1+vivienda_inflacionSeguro)^({year}-2026)*IF({aduB}=1,1+loc_precio_aduCosto/{basePr},1),0)'),
    ('ptax', 'Impuesto predial', '={val0}*loc_impuestoPropPct'),
    ('mant', 'Mantenimiento', '={val0}*vivienda_mantenimientoPct'),
    ('pmi', 'PMI', '=IF(AND({pmiOn0}=1,{mA}>0.8*{orig}),{orig}*(1-vivienda_enganchePct)*vivienda_pmiTasa,0)'),
    ('pmiOn', 'PMI activo (después)', '=IF(AND({pmiOn0}=1,{mA}>0.8*{orig}),1,0)'),
    ('rentI', 'Renta recibida', '=IF({own}=1,{units}*loc_rentaUnidad*12*{fI}*(1-vivienda_vacancia)*(1-vivienda_gestionPct),0)'),
    ('hous', 'Costo de vivienda', '=IF({own}=1,{mPaid}+{aPaid}+{ins}+{ptax}+{mant}+{pmi},loc_rentaMensual*12*{fI})'),
    ('fam', 'Aporte familiar', '=IF({own}=1,personal_aporteFamiliaMensual*12,0)'),
    # deudas
    ('irsP', 'IRS pagado', '=MIN(IF({k}=0,personal_irsSaldo,{irsE-1}),personal_irsPagoMensual*12)'),
    ('irsE', 'IRS saldo', '=IF({k}=0,personal_irsSaldo,{irsE-1})-{irsP}'),
    ('fedP', 'Federal inicio', '=IF({k}=0,IF(personal_deudaEsFederal=1,personal_deudaEstudiantil,0),{fedE-1})'),
    ('fRP', 'Tasa federal', '=IF({k}=0,personal_tasaDeudaEstudiantil,{fRE-1})'),
    ('fYP', 'Años repago fed.', '=IF({k}=0,0,{fYE-1})'),
    ('inS', 'En estudios (≥50%)', '=IF({study}>=0.5,1,0)'),
    ('useP', 'Usa RAP/PSLF', '=IF(AND({inS}=0,prestamos_pslfActivo=1,{pslf}>=0.5,{fedP}>0),1,0)'),
    ('fInt', 'Interés federal', '={fedP}*{fRP}'),
    ('rap', 'Pago RAP', '=MAX(prestamos_rapPagoMinimoAnual,IF({agi}<10000,0,MIN(0.1,INT({agi}/10000)/100))*{agi})'),
    ('fPay', 'Pago federal', '=IF({fedP}<=0,0,IF({inS}=1,0,IF({useP}=1,MIN({fedP}+{fInt},{rap}),PMT({fRP},MAX(1,prestamos_plazoAnios-{fYP}),-{fedP}))))'),
    ('fA', 'Federal tras pago', '=IF({fedP}<=0,0,IF({inS}=1,{fedP}+{fInt},IF({useP}=1,MAX(0,{fedP}-MAX({fPay}-{fInt},MIN(600,{fPay}))),MAX(0,{fedP}*(1+{fRP})-{fPay}))))'),
    ('psM', 'Meses PSLF', '=IF({k}=0,0,{psM-1})+IF({useP}=1,12*MIN(1,{pslf}),0)'),
    ('forg', 'Condonado PSLF', '=IF(AND({useP}=1,{psM}>=prestamos_pslfPagos,{fA}>0),{fA},0)'),
    ('fA2', 'Federal tras PSLF', '={fA}-{forg}'),
    ('fY1', 'Años repago (tras pago)', '=IF(AND({fedP}>0,{inS}=0,{useP}=0),{fYP}+1,{fYP})'),
    ('prvP', 'Privado inicio', '=IF({k}=0,IF(personal_deudaEsFederal=1,0,personal_deudaEstudiantil),{prvE-1})'),
    ('pRP', 'Tasa privada', '=IF({k}=0,IF(personal_deudaEsFederal=1,prestamos_tasaPrivada,personal_tasaDeudaEstudiantil),{pRE-1})'),
    ('pYP', 'Años repago priv.', '=IF({k}=0,0,{pYE-1})'),
    ('pPay', 'Pago privado', '=IF(OR({prvP}<=0,{inS}=1),0,PMT({pRP},MAX(1,prestamos_plazoAnios-{pYP}),-{prvP}))'),
    ('pA', 'Privado tras pago', '=IF({prvP}<=0,0,IF({inS}=1,{prvP}*(1+{pRP}),MAX(0,{prvP}*(1+{pRP})-{pPay})))'),
    ('pY1', 'Años repago priv. (tras)', '=IF(AND({prvP}>0,{inS}=0),{pYP}+1,{pYP})'),
    ('cA', 'Consumo tras interés', '=IF({k}=0,0,{conE-1})*(1+general_tasaDeudaConsumo)'),
    # flujo
    ('trX', 'Gasto extra travel', '={trav}*12*tiempos_travelGastoExtraMes*{fI}'),
    ('out', 'Salidas', '={liv}+{trX}+{hous}+{irsP}+{fPay}+{pPay}+{tuit}+{ev}-{rentI}-{fam}'),
    ('sur', 'Excedente (déficit)', '={take}-{out}'),
    ('cum', 'Flujo acumulado', '=IF({k}=0,0,{cum-1})+{sur}'),
    ('tgt', 'Meta enganche', '={price}*vivienda_enganchePct+{price}*vivienda_cierrePct+general_reservaMeses/12*{liv}'),
    ('lr', 'Retorno del ahorro líquido', '=IF({own}=1,g_ret,IF({L0}>0,MIN(1,{tgt}/{L0}),1)*g_cash+(1-IF({L0}>0,MIN(1,{tgt}/{L0}),1))*g_ret)'),
    ('pc', 'Pago deuda consumo', '=IF({sur}>=0,MIN({cA},{sur}),0)'),
    ('disc', 'Gasto discrecional', '=IF({sur}>=0,(1-general_tasaAhorro)*({sur}-{pc}),0)'),
    ('inv0', 'Ahorro', '=IF({sur}>=0,{sur}-{pc}-{disc},0)'),
    ('roth', 'Roth', '=MIN(personal_rothAnual,{inv0})'),
    ('prep', 'Prepago', '=IF(AND({sur}>=0,{own}=1,{mA}+{aA}>0),MIN(vivienda_prepagoPct*({inv0}-{roth}),{mA}+{aA}),0)'),
    ('toM', 'Prepago a hipoteca', '=MIN({prep},{mA})'),
    ('fromL', 'Retiro del ahorro', '=IF({sur}<0,MIN(MAX(0,{L0}),-{sur}),0)'),
    ('d2', 'Déficit por financiar', '=IF({sur}<0,-{sur}-{fromL},0)'),
    ('fAgP', 'Federal acumulado (antes)', '=IF({k}=0,0,{fAgE-1})'),
    ('nF', 'Nuevo préstamo federal', '=IF(AND({d2}>0,{study}>0),MIN({d2},MIN(prestamos_topeAnualPosgrado*MIN(1,{study}),MAX(0,prestamos_topeTotalPosgrado-{fAgP}))),0)'),
    ('nP', 'Nuevo préstamo privado', '=IF(AND({d2}>0,{study}>0),{d2}-{nF},0)'),
    ('nC', 'Nueva deuda consumo', '=IF(AND({d2}>0,{study}<=0),{d2},0)'),
    ('flow', 'Flujo al ahorro líquido', '=IF({sur}>=0,{inv0}-{roth}-{prep},-{fromL})'),
    ('liq', 'Ahorro líquido fin', '={L0}*(1+{lr})+{flow}*(1+{lr}/2)'),
    ('ret', 'Retiro fin', '=IF({k}=0,personal_retiroInicial,{ret-1})*(1+g_ret)+({k401}+{emp}+{roth})*(1+g_ret/2)'),
    ('valE', 'Valor vivienda fin', '=IF({own}=1,{val0}*(1+g_appr),0)'),
    ('mE', 'Hipoteca fin', '={mA}-{toM}'),
    ('aE', 'ADU fin', '={aA}-({prep}-{toM})'),
    ('fedE', 'Federal fin', '={fA2}+{nF}'),
    ('fRE', 'Tasa federal fin', '=IF({nF}>0,({fA2}*{fRP}+{nF}*prestamos_tasaFederalPosgrado)/({fA2}+{nF}),{fRP})'),
    ('fYE', 'Años repago fed. fin', '=IF({nF}>0,0,{fY1})'),
    ('fAgE', 'Federal acumulado', '={fAgP}+{nF}'),
    ('prvE', 'Privado fin', '={pA}+{nP}'),
    ('pRE', 'Tasa privada fin', '=IF({nP}>0,({pA}*{pRP}+{nP}*prestamos_tasaPrivada)/({pA}+{nP}),{pRP})'),
    ('pYE', 'Años repago priv. fin', '=IF({nP}>0,0,{pY1})'),
    ('conE', 'Deuda consumo fin', '={cA}-{pc}+{nC}'),
    ('payY', 'Año vivienda pagada', '=IF(IF({k}=0,"",{payY-1})<>"",{payY-1},IF(AND({own}=1,{mE}<=0.5,{aE}<=0.5,OR({isAdu}=0,{aduB}=1)),{year},""))'),
    ('debt', 'Deuda total', '={mE}+{aE}+{fedE}+{prvE}+{conE}+{irsE}'),
    ('nw', 'PATRIMONIO NETO', '={liq}+{ret}+{valE}-{debt}'),
    ('nwR', 'Patrimonio (dólares 2026)', '={nw}/(1+g_inf)^({year}+1-2026)'),
]
SIMCOL = {key: i + 1 for i, (key, _, _) in enumerate(SIM)}
NWCOL = L(SIMCOL['nw'])


def render_formula(tpl, r, k):
    def rep(m):
        tok = m.group(1)
        if tok == 'OV':
            return f'${L(OVC + k)}$5:${L(OVC + k)}$29'
        if tok == 'k':
            return str(k)
        if tok == 'r':
            return str(r)
        if tok.endswith('-1'):
            return f'{L(SIMCOL[tok[:-2]])}{r - 1}'
        return f'{L(SIMCOL[tok])}{r}'
    out = tpl.replace('{{', '\x00').replace('}}', '\x01')
    out = re.sub(r'\{([A-Za-z0-9]+(?:-1)?)\}', rep, out)
    return out.replace('\x00', '{').replace('\x01', '}')


SHEETNAMES = {}
for p in D['paths']:
    title = p['nombre'].split(' · ')[0].replace('★', 'x') + ' ' + p['id']
    title = re.sub(r'[\[\]\*\?/\\:]', '', title)[:31]
    SHEETNAMES[p['id']] = title
    ws = wb.create_sheet(title)
    ws['A1'] = p['nombre']; ws['A1'].font = TITLE
    ws['A2'] = 'Etapas (celdas azules editables; fechas decimales: 2027.67 = sep-2027). Traslape por año en columnas T→AM. Hitos con costo en AN→AP.'
    ws['A2'].font = NOTE
    hdr(ws, 4, PH_COLS)
    phases = p['phases']
    assert len(phases) <= MAXPH, p['id']
    for i in range(MAXPH):
        rr = 5 + i
        if i < len(phases):
            ph = phases[i]
            role = ph.get('role')
            vals = [ph['lbl'], round(ph['from'], 6), round(ph['to'], 6), role or ph.get('scorp') or ph.get('se') or '']
            if role:
                sal, rai = '=' + role_salary_expr(role), '=' + raise_expr(role)
            elif ph.get('scorp') == 'locums':
                sal, rai = 0, '=roles_hospitalista_aumento'
            elif ph.get('scorp') or ph.get('se'):
                sal, rai = 0, '=g_inf'
            else:
                sal, rai = 0, 0
            bonus = ph.get('bonus') or 0
            bonus = '=tiempos_bonoFirmaICU' if role == 'icuNuevo' and bonus else ('=md_bonoFirmaAttending' if role in ('hospitalista', 'fm') and bonus else bonus)
            scorp = '=md_locumsTurnosMes*12*12*md_locumsTarifa*g_salMult' if ph.get('scorp') == 'locums' else ('=md_directorMedico' if ph.get('scorp') == 'director' else 0)
            se = '=md_tutoria' if ph.get('se') == 'tutoria' else 0
            vals += [sal, rai, hours_expr(ph), mult_expr(ph), bonus, tuition_expr(ph), 1 if ph.get('study') else 0,
                     1 if ph.get('pslf') else 0, 1 if ph.get('match') else 0, 1 if role == 'travel' else 0,
                     '=tiempos_travelStipendPct' if role == 'travel' else 0, scorp, se, 1 if ph.get('attending') else 0]
            if role == 'crna' and ph.get('pslf') is not None:
                vals[11] = '=tiempos_crnaPslf'
            if role in ('residente', 'hospitalista', 'fm'):
                vals[11] = '=md_residentePslf'
        else:
            vals = ['', 0, 0, '', 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]
        for j, v in enumerate(vals):
            c = ws.cell(row=rr, column=1 + j, value=v)
            if j in (1, 2) or (isinstance(v, (int, float)) and j >= 6):
                c.font = INPUT
    # matriz de traslape
    for k, y in enumerate(YEARS):
        c = ws.cell(row=4, column=OVC + k, value=y); c.font = BOLD; c.fill = HDR_FILL
        for i in range(MAXPH):
            rr = 5 + i
            ws.cell(row=rr, column=OVC + k, value=f'=MAX(0,MIN($C{rr},{L(OVC + k)}$4+1)-MAX($B{rr},{L(OVC + k)}$4))')
    # hitos
    hdr(ws, 4, ['Hito', 'Fecha', 'Costo (2026$)'], col=EVC + 1)
    evs = [e for e in p['events']]
    assert len(evs) <= MAXEV, p['id']
    for i in range(MAXEV):
        rr = 5 + i
        if i < len(evs):
            e = evs[i]
            ws.cell(row=rr, column=EVC + 1, value=e['lbl'])
            c = ws.cell(row=rr, column=EVC + 2, value=round(e['at'], 6)); c.font = INPUT
            c = ws.cell(row=rr, column=EVC + 3, value=event_cost_expr(e)); c.font = INPUT
        else:
            ws.cell(row=rr, column=EVC + 2, value=0); ws.cell(row=rr, column=EVC + 3, value=0)
    assert L(EVC + 2) == 'AO' and L(EVC + 3) == 'AP', (L(EVC + 2), L(EVC + 3))
    # simulación anual
    ws.cell(row=Y0 - 2, column=1, value='Simulación anual (todas las celdas son fórmulas)').font = TITLE
    hdr(ws, Y0 - 1, [h for _, h, _ in SIM])
    for k, y in enumerate(YEARS):
        r = Y0 + k
        for key, _, tpl in SIM:
            col = SIMCOL[key]
            if key == 'year':
                ws.cell(row=r, column=col, value=y)
            elif key == 'k':
                ws.cell(row=r, column=col, value=k)
            else:
                ws.cell(row=r, column=col, value=render_formula(tpl, r, k)).number_format = '#,##0'
    ws.freeze_panes = ws.cell(row=Y0, column=2)
    ws.column_dimensions['A'].width = 40
    nwc = ws.cell(row=Y0 - 1, column=SIMCOL['nw']); nwc.fill = PatternFill('solid', fgColor='FDE68A')

# ----------------------------------------------------------------------------- Comparativa
C = wb.create_sheet('Comparativa', 1)
C['A1'] = 'Comparativa de caminos (ubicación, vivienda y escenario según la hoja Supuestos)'; C['A1'].font = TITLE
C['A2'] = '=CONCATENATE("Ubicación: ",sel_loc,"  ·  Vivienda: ",sel_cfg,"  ·  Escenario: ",Supuestos!B5)'
hdr(C, 4, ['Año'] + [p['nombre'] for p in D['paths']])
for k, y in enumerate(YEARS):
    C.cell(row=5 + k, column=1, value=y)
    for j, p in enumerate(D['paths']):
        C.cell(row=5 + k, column=2 + j, value=f"='{SHEETNAMES[p['id']]}'!{NWCOL}{Y0 + k}").number_format = '#,##0'
last = 5 + H - 1
sr = last + 2
labels = [('Patrimonio año 10', lambda col: f'={col}{5 + 9}'), ('Patrimonio año 20', lambda col: f'={col}{last}'),
          ('Año compra vivienda', None), ('Año vivienda pagada', None), ('Deuda máxima', None), ('Ranking (año 20)', lambda col: f'=RANK({col}{last},$B${last}:${L(1 + len(D["paths"]))}${last})')]
for i, (lbl, fn) in enumerate(labels):
    C.cell(row=sr + i, column=1, value=lbl).font = BOLD
    for j, p in enumerate(D['paths']):
        col = L(2 + j)
        sh = f"'{SHEETNAMES[p['id']]}'"
        if lbl == 'Año compra vivienda':
            f = f'=IF({sh}!{L(SIMCOL["buyY"])}{Y0 + H - 1}="","no viable",{sh}!{L(SIMCOL["buyY"])}{Y0 + H - 1})'
        elif lbl == 'Año vivienda pagada':
            f = f'=IF({sh}!{L(SIMCOL["payY"])}{Y0 + H - 1}="","—",{sh}!{L(SIMCOL["payY"])}{Y0 + H - 1})'
        elif lbl == 'Deuda máxima':
            f = f'=MAX({sh}!{L(SIMCOL["debt"])}{Y0}:{L(SIMCOL["debt"])}{Y0 + H - 1})'
        else:
            f = fn(col)
        c = C.cell(row=sr + i, column=2 + j, value=f)
        if lbl not in ('Año compra vivienda', 'Año vivienda pagada', 'Ranking (año 20)'):
            c.number_format = '#,##0'
# matriz de cruces
mr = sr + len(labels) + 2
C.cell(row=mr, column=1, value='Cruces: año desde el cual el camino de la FILA supera permanentemente al de la COLUMNA').font = BOLD
hdr(C, mr + 1, ['Camino'] + [p['nombre'].split(' · ')[0] for p in D['paths']])
n = len(D['paths'])
for i, pi in enumerate(D['paths']):
    rr = mr + 2 + i
    C.cell(row=rr, column=1, value=pi['nombre'])
    ci = L(2 + i)
    for j in range(n):
        if i == j:
            C.cell(row=rr, column=2 + j, value='—'); continue
        cj = L(2 + j)
        f = (f'=IF(SUMPRODUCT(--({ci}$5:{ci}${last}<={cj}$5:{cj}${last}))=0,"siempre",'
             f'IF(LOOKUP(2,1/({ci}$5:{ci}${last}<={cj}$5:{cj}${last}),$A$5:$A${last})>={YEARS[-1]},"nunca",'
             f'LOOKUP(2,1/({ci}$5:{ci}${last}<={cj}$5:{cj}${last}),$A$5:$A${last})+1))')
        C.cell(row=rr, column=2 + j, value=f)
C.column_dimensions['A'].width = 42
for j in range(n):
    C.column_dimensions[L(2 + j)].width = 16
C.freeze_panes = 'B5'
nr = mr + 3 + n
C.cell(row=nr, column=1, value='Notas').font = BOLD
notes = [
    'El Excel replica la lógica del dashboard (src/engine.cjs). El test tests/test_excel.py recalcula este libro con LibreOffice y compara el patrimonio con el motor.',
    'Las fechas de las etapas (hojas de cada camino) son valores; el escenario "retraso" del dashboard no las mueve. El Monte Carlo del camino MD solo está en el dashboard.',
    'Para cambiar ubicación, vivienda o escenario usa los selectores en Supuestos!B3:B5.',
]
for i, t in enumerate(notes):
    C.cell(row=nr + 1 + i, column=1, value=t).font = NOTE

wb.calculation.fullCalcOnLoad = True
out = os.path.join(ROOT, 'modelo_carrera.xlsx')
wb.save(out)
print('Generado', out, '·', len(wb.sheetnames), 'hojas ·', len(names), 'nombres definidos')
