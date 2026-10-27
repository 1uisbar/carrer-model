#!/usr/bin/env python3
"""Paridad Excel ↔ motor JS: regenera el libro, lo recalcula con LibreOffice y compara el patrimonio
año a año de todos los caminos (ubicación y escenario por defecto + variantes)."""
import json, os, shutil, subprocess, sys, tempfile
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOL = 0.005  # 0.5% del patrimonio (o $500)


def recalc(src, loc, cfg, esc, tmp):
    wb = openpyxl.load_workbook(src)
    S = wb['Supuestos']
    S['B3'], S['B4'], S['B5'] = loc, cfg, esc
    p = os.path.join(tmp, f'in_{loc}_{cfg}_{esc}.xlsx')
    wb.save(p)
    out = os.path.join(tmp, 'out')
    os.makedirs(out, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out, p], check=True, capture_output=True, timeout=600)
    return openpyxl.load_workbook(os.path.join(out, os.path.basename(p)), data_only=True)


def main():
    subprocess.run([sys.executable, os.path.join(ROOT, 'excel', 'build_excel.py')], check=True, capture_output=True)
    src = os.path.join(ROOT, 'modelo_carrera.xlsx')
    d = json.load(open(os.path.join(ROOT, 'excel', 'paths_default.json')))
    cases = [('dfw', 'duplex', 'base'), ('miami', 'duplex', 'pesimista'), ('cleveland', 'duplex', 'optimista'), ('nyc', 'duplex', 'base'), ('houston', 'duplex', 'base'), ('dfw', 'casaAdu', 'base'), ('dfw', 'multi', 'base'), ('dfw', 'inlaw', 'base')]
    tmp = tempfile.mkdtemp()
    fails = 0
    try:
        for loc, cfg, esc in cases:
            wb = recalc(src, loc, cfg, esc, tmp)
            C = wb['Comparativa']
            res = d['results'][f'{loc}|{esc}' + ('' if cfg == 'duplex' else '|' + cfg)]
            A = json.load(open(os.path.join(ROOT, 'supuestos.json')))
            delay = A['escenarios'][esc]['retraso'] != 0
            for j, p in enumerate(d['paths']):
                # El 'retraso' del escenario mueve fechas en el motor; en el Excel las fechas son fijas (documentado).
                if delay and (p['id'].startswith('md_') or 'crna' in p['id'] or p['id'] in ('fnp', 'pmhnp')):
                    continue
                xs = [C.cell(row=5 + k, column=2 + j).value for k in range(20)]
                js = res[p['id']]['nw']
                bad = [(2027 + k, x, y) for k, (x, y) in enumerate(zip(xs, js)) if not isinstance(x, (int, float)) or abs(x - y) > max(500, TOL * abs(y))]
                if bad:
                    fails += 1
                    print(f'FALLA {loc}/{esc}/{p["id"]}: {bad[:3]}')
            print(f'revisado {loc}/{cfg}/{esc} (fallas acumuladas: {fails})')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if fails:
        sys.exit(1)
    print('Paridad Excel ↔ motor: OK')


if __name__ == '__main__':
    main()
